from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\nosql\project\RankRush\rankrush-realtime")
OUT = ROOT / "output" / "Bao_cao_ky_thuat_va_luong_xu_ly_RankRush.docx"

NAVY = "163A4A"
BLUE = "218FAC"
CYAN = "2BA5C3"
DARK = "17202A"
MUTED = "667085"
LIGHT = "E8EEF5"
LIGHTER = "F4F6F9"
GREEN = "18794E"
GOLD = "9A6700"
RED = "B42318"
WHITE = "FFFFFF"
MONO = "Consolas"
FONT = "Calibri"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_repeat_table_header(row):
    repeat_header(row)


def set_table_geometry(table, widths: Sequence[int], indent=TABLE_INDENT):
    assert sum(widths) == TABLE_WIDTH, (widths, sum(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        no_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_external_hyperlink(paragraph, text: str, target: str, color=BLUE):
    part = paragraph.part
    rid = part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), MONO)
    rfonts.set(qn("w:hAnsi"), MONO)
    rpr.append(rfonts)
    clr = OxmlElement("w:color")
    clr.set(qn("w:val"), color)
    rpr.append(clr)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    link.append(new_run)
    paragraph._p.append(link)


def file_url(rel: str) -> str:
    return (ROOT / rel).resolve().as_uri()


def add_code_ref(doc, rel: str, lines: str, note: str = ""):
    p = doc.add_paragraph()
    p.style = doc.styles["Code Reference"]
    p.add_run("Mã nguồn: ")
    # Keep references as deterministic local paths + line numbers. External file
    # hyperlinks make some Windows Word builds spend minutes resolving links
    # while paginating/exporting long reports.
    r = p.add_run(f"{rel}:{lines}")
    set_run_font(r, name=MONO, size=8.5, color=BLUE, bold=True)
    if note:
        p.add_run(f" — {note}")
    return p


def add_para(doc, text: str = "", bold=False, italic=False, color=DARK, size=11, align=None,
             before=0, after=6, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep_with_next
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=11, color=DARK)
    return p


def add_number(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=11, color=DARK)
    return p


def add_heading(doc, text: str, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_callout(doc, title: str, body: str, fill=LIGHTER, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r = p2.add_run(body)
    set_run_font(r, size=10.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int],
              header_fill=LIGHT, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        shade(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=DARK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code_block(doc, code: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    shade(cell, "F7F9FB")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(code.rstrip().splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name=MONO, size=8.2, color="263238")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, "1F4D78", 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    code_ref = styles.add_style("Code Reference", 1)
    code_ref.font.name = MONO
    code_ref._element.rPr.rFonts.set(qn("w:ascii"), MONO)
    code_ref._element.rPr.rFonts.set(qn("w:hAnsi"), MONO)
    code_ref.font.size = Pt(8.5)
    code_ref.font.color.rgb = rgb(MUTED)
    code_ref.paragraph_format.space_before = Pt(2)
    code_ref.paragraph_format.space_after = Pt(5)
    code_ref.paragraph_format.keep_with_next = False


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("RANKRUSH  |  BÁO CÁO KỸ THUẬT & LUỒNG XỬ LÝ")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    p = footer.paragraphs[0]
    add_page_field(p)


def add_toc(doc):
    add_heading(doc, "Mục lục nội dung", 1)
    items = [
        "1. Phạm vi, phương pháp và kết luận nhanh",
        "2. Kiến trúc tổng thể và công nghệ sử dụng",
        "3. Phân tích lựa chọn cơ sở dữ liệu",
        "4. Thiết kế dữ liệu Redis và cách đọc/ghi",
        "5. TTL — thời gian sống và vòng đời phiên",
        "6. Bảo mật, xác thực và cách nhận biết đúng người dùng",
        "7. Luồng chi tiết từng nhóm chức năng",
        "8. Trí tuệ nhân tạo, PDF/CSV và kiểm định câu hỏi",
        "9. Quản trị, sao lưu, giám sát và vận hành",
        "10. Tối ưu, giới hạn và lộ trình nâng cấp",
        "11. Danh mục API và chỉ mục mã nguồn",
    ]
    for item in items:
        add_para(doc, item, size=10.5, color=NAVY, after=4)
    add_callout(
        doc,
        "Cách dùng tham chiếu mã nguồn",
        "Mỗi mục có dòng “Mã nguồn: tệp:Lx-Ly”. Dùng đường dẫn và số dòng này để mở đúng vị trí trong mã nguồn; "
        "số dòng chỉ vị trí đã kiểm tra tại nhánh develop, commit 5932792 ngày 02/08/2026.",
    )


def build():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)

    # Cover: editorial_cover pattern, compact_reference_guide tokens.
    add_para(doc, "BÁO CÁO PHÂN TÍCH HỆ THỐNG", size=11, color=CYAN, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=88, after=18)
    add_para(doc, "RANKRUSH", size=32, color=NAVY, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    add_para(doc, "Công nghệ, cơ sở dữ liệu Redis và toàn bộ luồng xử lý chức năng",
             size=16, color=BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    add_para(
        doc,
        "Báo cáo kỹ thuật có truy vết trực tiếp tới mã nguồn",
        size=11.5,
        color=MUTED,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=72,
    )
    add_table(
        doc,
        ["Thuộc tính", "Giá trị"],
        [
            ["Dự án", "rankrush-realtime"],
            ["Phạm vi rà soát", "Frontend, API, Redis, realtime, AI, email, quản trị, Docker và kiểm thử"],
            ["Nhánh / commit", "develop / 5932792"],
            ["Ngày lập", "02/08/2026 (Asia/Saigon)"],
            ["Đường dẫn nguồn", str(ROOT)],
        ],
        [2700, 6660],
        header_fill="DDECF1",
        font_size=9.5,
    )
    add_para(
        doc,
        "Lưu ý bảo mật: báo cáo chỉ nêu tên biến môi trường; không sao chép SMTP_PASS, GEMINI_API_KEY "
        "hoặc giá trị bí mật từ tệp .env.",
        size=9.5,
        color=RED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=0,
    )
    doc.add_page_break()
    add_toc(doc)
    doc.add_page_break()

    # 1
    add_heading(doc, "1. Phạm vi, phương pháp và kết luận nhanh", 1)
    add_para(
        doc,
        "Báo cáo được xây dựng bằng cách đọc cấu hình workspace, package manifests, API routes, "
        "các cấu trúc Redis, mã frontend và các luồng Socket.IO. Mọi kết luận quan trọng đều gắn "
        "với tệp và số dòng. “Toàn bộ chức năng” được tổ chức theo nhóm hành vi người dùng và endpoint, "
        "thay vì chép lại từng dòng mã.",
    )
    add_heading(doc, "1.1 Kết luận kiến trúc", 2)
    for text in [
        "Hệ thống là monorepo npm gồm web React/Vite và API Node.js/Express/Socket.IO.",
        "Cơ sở dữ liệu ứng dụng hiện tại là Redis 7.4; không có driver, schema hay kết nối MongoDB trong mã nguồn.",
        "Redis được dùng vừa làm kho dữ liệu chính, vừa làm cache TTL, bảng xếp hạng ZSET, nhật ký Stream, khóa phân tán và rate limiter.",
        "Mật khẩu tài khoản thường được lưu dưới dạng bcrypt hash; mã email được lưu dưới SHA-256 digest có TTL. JWT xác định host/admin/player.",
        "Phiên game có TTL mặc định 7 ngày; PIN tham gia 1 ngày; mã đăng ký 10 phút; mã đặt lại 15 phút. Các giới hạn đều điều chỉnh được trong .env.",
        "Thiết kế rất phù hợp demo và game realtime quy mô vừa; sản xuất dài hạn nên bổ sung bảo mật Redis, giám sát bộ nhớ, HA và có thể kho SQL bền vững cho phân tích/audit.",
    ]:
        add_bullet(doc, text)
    add_code_ref(doc, "package.json", "L5-L23", "workspace, scripts và Node >= 20")
    add_code_ref(doc, "docker-compose.yml", "L1-L28", "Redis 7.4, AOF và RedisInsight")
    add_code_ref(doc, "apps/api/src/config.ts", "L11-L64", "cấu hình runtime và tiền tố key")

    add_heading(doc, "1.2 Phạm vi chức năng được phân tích", 2)
    add_table(
        doc,
        ["Nhóm", "Chức năng"],
        [
            ["Tài khoản", "Đăng ký + email OTP, đăng nhập, quên/đặt lại mật khẩu, hồ sơ, avatar, phân quyền"],
            ["Quiz", "Tạo, sửa, xóa, clone, xuất bản, thư viện, tìm kiếm, luyện tập, 7 dạng câu hỏi"],
            ["AI/import", "Sinh từ chủ đề/PDF, nhập CSV, Gemini/Ollama/local fallback, review và chống trùng"],
            ["Game", "Tạo phòng/PIN, QR, join, lobby, start, phase timer, trả lời, điểm, leaderboard, pause/resume/skip/end/replay"],
            ["Realtime", "Socket auth, room host/player, snapshot riêng theo quyền, online/offline, tự kết thúc phòng rỗng"],
            ["Admin", "Tổng quan, người dùng, phòng, hoạt động, Redis/system, sao lưu/khôi phục"],
            ["Trải nghiệm", "VI/EN, sáng/tối, nhạc, voice đọc câu hỏi, copy link/PIN, responsive UI"],
        ],
        [1700, 7660],
    )

    # 2
    doc.add_page_break()
    add_heading(doc, "2. Kiến trúc tổng thể và công nghệ sử dụng", 1)
    add_code_block(
        doc,
        """Trình duyệt (React + TypeScript + Vite)
  ├─ HTTP/JSON + Bearer JWT ───────────────┐
  └─ Socket.IO realtime ───────────────────┤
                                           ▼
API Node.js + Express + Socket.IO
  ├─ Xác thực: Zod → bcrypt → JWT
  ├─ Nghiệp vụ: quiz, session, scoring, admin
  ├─ Email: Nodemailer → SMTP Gmail
  ├─ AI: Gemini → Ollama → local fallback
  └─ Persistence: ioredis
                                           ▼
Redis 7.4
  Hash · String · Set · List · ZSET · Stream · TTL · Lua
                                           ▼
Docker volume + AOF everysec; RedisInsight tại cổng 5540""",
    )
    add_code_ref(doc, "apps/api/src/index.ts", "L1-L169", "khởi tạo HTTP, Express, Socket.IO, middleware và upload")
    add_code_ref(doc, "apps/web/src/App.tsx", "L6905-L6926", "toàn bộ router giao diện")

    technologies = [
        ("TypeScript", "Ngôn ngữ JavaScript có kiểu tĩnh, biên dịch về JS.", "Phát hiện sai kiểu sớm; dùng chung type; IDE tốt.", "Không bảo đảm dữ liệu runtime; cần build.", "Cả frontend/backend cùng ngôn ngữ, giảm lệch hợp đồng API."),
        ("Node.js 20+", "Runtime JavaScript bất đồng bộ, event loop.", "I/O đồng thời tốt; hệ sinh thái npm; phù hợp socket.", "CPU-heavy có thể chặn event loop.", "Game chủ yếu I/O mạng/Redis và realtime."),
        ("React 18", "UI khai báo theo component và state.", "Tái sử dụng, state-driven, hệ sinh thái lớn.", "App.tsx quá lớn sẽ khó bảo trì nếu không tách.", "Nhiều màn hình động, timer, form, socket snapshot."),
        ("Vite 6", "Dev server/build tool dựa ESM.", "Khởi động và HMR nhanh; cấu hình gọn.", "Không thay thế backend; phụ thuộc trình duyệt hiện đại.", "Tăng tốc vòng lặp phát triển UI."),
        ("React Router 7", "Định tuyến SPA phía client.", "URL rõ, điều hướng không reload.", "Cần guard quyền ở cả client và server.", "Nhiều trang auth/editor/host/player/admin."),
        ("Express 4", "Web framework middleware cho Node.", "Đơn giản, linh hoạt, thư viện nhiều.", "Dễ tạo tệp route quá lớn; ít quy ước kiến trúc.", "API hiện có quy mô vừa, cần phát triển nhanh."),
        ("Socket.IO 4", "Kênh event realtime có room và reconnect.", "Broadcast theo phòng, tự reconnect, fallback transport.", "Tốn state/kết nối; scale cần adapter chia sẻ.", "Đồng bộ lobby, câu hỏi, tiến độ và leaderboard tức thời."),
        ("Redis 7.4", "Kho key-value in-memory đa cấu trúc.", "Độ trễ thấp, TTL, ZSET, Stream, atomic Lua.", "RAM đắt; truy vấn ad-hoc/quan hệ yếu.", "Một game cần state nhanh, rank trực tiếp và TTL tự dọn."),
        ("ioredis", "Client Redis cho Node.", "Pipeline, MULTI, Lua, reconnect.", "Phụ thuộc Redis; cần xử lý lỗi mạng.", "Hỗ trợ đầy đủ lệnh mà store.ts sử dụng."),
        ("Zod", "Xác thực schema ở runtime.", "Biên dữ liệu rõ; lỗi validation nhất quán.", "Viết schema lặp với type nếu không sinh tự động.", "Không tin dữ liệu form/client/API."),
        ("bcryptjs", "KDF chậm có salt để băm mật khẩu.", "Khó brute-force hơn hash nhanh.", "Tốn CPU; cần cost phù hợp.", "Đúng mục đích lưu mật khẩu, cost 12."),
        ("JWT", "Token ký chứa claims và hạn dùng.", "API stateless; phân biệt host/admin/player.", "Khó thu hồi tức thì nếu không có deny-list.", "Client HTTP và Socket cùng xác thực bằng token."),
        ("Nodemailer/SMTP", "Gửi email qua SMTP.", "Phổ biến; hỗ trợ Gmail App Password.", "Phụ thuộc cấu hình/nhà cung cấp, có quota.", "OTP đăng ký và quên mật khẩu gửi trực tiếp."),
        ("Multer + pdf-parse", "Nhận multipart in-memory và trích text PDF.", "Luồng upload đơn giản; không cần tệp tạm.", "PDF scan ảnh không OCR; RAM tăng theo file.", "Nguồn sinh quiz từ PDF tối đa 10 MB."),
        ("Gemini SDK", "LLM cloud sinh JSON có schema và review.", "Chất lượng/ngữ cảnh cao; structured output.", "Chi phí, mạng, key và quota.", "Cải thiện câu hỏi so với local model."),
        ("Ollama", "Chạy LLM cục bộ qua HTTP.", "Riêng tư, offline, không phí theo request.", "Chất lượng/tốc độ phụ thuộc máy và model.", "Fallback qwen2.5:3b khi Gemini không dùng được."),
        ("Docker Compose", "Khai báo dịch vụ container.", "Môi trường Redis tái lập; volume/publish port.", "Cần Docker Desktop; tài nguyên nền.", "Khởi tạo Redis/RedisInsight nhất quán."),
        ("RedisInsight", "GUI quan sát Redis.", "Xem key/type/TTL/command thuận tiện.", "Không thay thế monitoring production.", "Hỗ trợ học tập, debug và trình bày CSDL."),
        ("Vitest/Supertest", "Unit/integration test cho TS và HTTP.", "Chạy nhanh; kiểm tra regression.", "Không tự bao phủ E2E/trình duyệt thật.", "Bảo vệ auth, score, import, backup, preferences."),
        ("Git/GitHub + VS Code", "Quản lý phiên bản và môi trường phát triển.", "Lịch sử, nhánh, review, IDE TypeScript.", "Cần quy trình commit/review tốt.", "Phù hợp dự án nhóm và truy vết thay đổi."),
    ]
    add_table(
        doc,
        ["Công nghệ", "Đặc điểm", "Ưu điểm", "Nhược điểm", "Tại sao chọn"],
        technologies,
        [1250, 1900, 1850, 1750, 2610],
        font_size=7.7,
    )
    add_code_ref(doc, "apps/api/package.json", "L16-L44", "dependencies và công cụ backend")
    add_code_ref(doc, "apps/web/package.json", "L12-L26", "dependencies và công cụ frontend")

    # 3
    doc.add_page_break()
    add_heading(doc, "3. Phân tích yêu cầu thực tế và lựa chọn cơ sở dữ liệu", 1)
    add_heading(doc, "3.1 Yêu cầu dữ liệu thực tế", 2)
    for text in [
        "Ghi/đọc nhanh theo session trong thời gian câu hỏi chỉ vài giây.",
        "Tăng điểm và lấy top/rank đồng thời cho nhiều người chơi.",
        "Ngăn trả lời lặp và join trùng một cách nguyên tử.",
        "Tự hết hạn PIN, OTP, rate-limit và dữ liệu phiên.",
        "Phát nhật ký sự kiện có thứ tự, giới hạn độ dài.",
        "Truy cập chính theo khóa đã biết: userId, quizId, sessionId, PIN; ít join quan hệ phức tạp.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "3.2 Ma trận lựa chọn", 2)
    add_table(
        doc,
        ["Tiêu chí", "Redis", "MongoDB", "PostgreSQL"],
        [
            ["Latency realtime", "Rất tốt: in-memory", "Tốt nhưng thường cao hơn", "Tốt, cần cache/index"],
            ["Leaderboard", "ZSET native: ZINCRBY/ZREVRANK", "Phải update + sort/index", "UPDATE + ORDER BY/window"],
            ["TTL", "Native trên key", "TTL index theo document", "Job/partition hoặc extension"],
            ["Atomic nhiều bước", "Lua/MULTI rất phù hợp cùng key-slot", "Transaction hỗ trợ nhưng nặng hơn", "ACID mạnh nhất"],
            ["Quan hệ/truy vấn ad-hoc", "Yếu", "Khá linh hoạt document", "Mạnh nhất"],
            ["Chi phí lưu dài hạn", "RAM cao", "Disk hiệu quả hơn", "Disk hiệu quả hơn"],
            ["Phù hợp hiện trạng", "Cao", "Trung bình", "Trung bình/cao nếu hybrid"],
        ],
        [1900, 2480, 2480, 2500],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Kết luận lựa chọn",
        "Redis là lựa chọn hợp lý cho lõi game realtime vì dữ liệu truy cập theo key, leaderboard và TTL là yêu cầu trung tâm. "
        "MongoDB không xuất hiện trong dự án hiện tại. Nếu cần báo cáo lịch sử phức tạp, giao dịch nghiệp vụ dài hạn hoặc audit "
        "không được phép mất, kiến trúc hybrid Redis + PostgreSQL là hướng nâng cấp an toàn hơn thay vì thay Redis hoàn toàn.",
        fill="EAF6F0",
        accent=GREEN,
    )

    add_heading(doc, "3.3 Tại sao không phải MongoDB trong bản hiện tại?", 2)
    add_para(
        doc,
        "Không có package mongodb/mongoose, URI MongoDB, model/schema hoặc lệnh MongoDB trong package manifests và mã API. "
        "Các thao tác lưu đều đi qua ioredis. MongoDB có ưu thế document và truy vấn linh hoạt, nhưng không cho một lệnh "
        "native tương đương ZINCRBY + ZREVRANK để cập nhật và đọc hạng với mô hình đơn giản như ZSET.",
    )
    add_code_ref(doc, "apps/api/package.json", "L16-L31", "có ioredis, không có MongoDB driver")
    add_code_ref(doc, "apps/api/src/redis.ts", "L1-L5", "client CSDL duy nhất")

    # 4
    add_heading(doc, "4. Thiết kế dữ liệu Redis và cách đọc/ghi", 1)
    add_heading(doc, "4.1 Kết nối Redis", 2)
    add_code_block(
        doc,
        """export const redis = new Redis(config.REDIS_URL, {
  maxRetriesPerRequest: 2,
  enableReadyCheck: true,
  lazyConnect: true
});
export async function connectRedis() {
  if (redis.status === "wait") await redis.connect();
  await redis.ping();
}""",
    )
    add_para(
        doc,
        "API đọc REDIS_URL và REDIS_PREFIX từ môi trường, tạo client lazy-connect, PING trước khi HTTP server listen. "
        "Mọi key có dạng rankrush:<phần...> để cô lập namespace. Khi dừng tiến trình, API QUIT Redis có kiểm soát.",
    )
    add_code_ref(doc, "apps/api/src/config.ts", "L6-L20, L62-L64")
    add_code_ref(doc, "apps/api/src/redis.ts", "L1-L5")
    add_code_ref(doc, "apps/api/src/index.ts", "L2977-L2987")

    add_heading(doc, "4.2 Danh mục key và cấu trúc dữ liệu", 2)
    key_rows = [
        ("rankrush:users", "SET", "Danh sách userId", "SADD / SMEMBERS", "Không TTL"),
        ("rankrush:user:<id>", "HASH", "Hồ sơ, role, status, passwordHash", "HSET / HGETALL", "Không TTL"),
        ("rankrush:user-email:<email>", "STRING", "Email → userId", "SET / GET", "Không TTL"),
        ("rankrush:user-username:<name>", "STRING", "Username → userId", "SET / GET", "Không TTL"),
        ("rankrush:email-verification:<email>", "STRING", "Digest OTP đăng ký", "SET EX / GET / DEL", "10 phút"),
        ("rankrush:password-reset:<email>", "STRING", "Digest OTP reset", "SET EX / GET / DEL", "15 phút"),
        ("rankrush:login-attempts:<id>", "STRING", "Đếm login sai", "INCR / GET", "15 phút"),
        ("rankrush:quizzes", "SET", "Danh sách quizId", "SADD / SMEMBERS", "Không TTL"),
        ("rankrush:quizzes:public", "ZSET", "Quiz public theo updatedAt", "ZADD / ZREVRANGE", "Không TTL"),
        ("rankrush:user:<id>:quizzes", "SET", "Quiz thuộc user", "SADD / SMEMBERS", "Không TTL"),
        ("rankrush:quiz:<id>", "HASH", "Metadata quiz", "HSET / HGETALL", "Không TTL"),
        ("rankrush:quiz:<id>:questions", "LIST", "Thứ tự questionId", "RPUSH/LRANGE/LREM", "Không TTL"),
        ("rankrush:question:<id>", "HASH", "Nội dung + đáp án + điểm", "HSET / HGETALL", "Không TTL"),
        ("rankrush:sessions", "SET", "Danh sách sessionId", "SADD / SMEMBERS", "Tự dọn id mồ côi"),
        ("rankrush:session-pin:<pin>", "STRING", "PIN → sessionId", "SET EX NX / GET", "1 ngày"),
        ("rankrush:session:{sid}", "HASH", "State/phase/settings/vòng đời", "HSET / HGETALL", "7 ngày"),
        ("rankrush:session:{sid}:players", "SET", "Danh sách playerId", "SADD / SMEMBERS", "7 ngày khi tồn tại"),
        ("rankrush:player:{sid}:<pid>", "HASH", "Nickname/avatar/team/online", "HSET / HGETALL", "7 ngày"),
        ("rankrush:leaderboard:{sid}", "ZSET", "playerId → điểm", "ZADD/ZINCRBY/ZREVRANGE/ZREVRANK", "7 ngày"),
        ("rankrush:teamboard:{sid}", "ZSET", "team → điểm", "ZADD/ZINCRBY/ZREVRANGE", "7 ngày"),
        ("rankrush:answer:{sid}:<pid>:<qid>", "HASH", "Một câu trả lời duy nhất", "HSET / EXISTS / HGETALL", "7 ngày"),
        ("rankrush:session:{sid}:answers", "SET", "Danh sách answer-key", "SADD / SMEMBERS", "7 ngày"),
        ("rankrush:events:{sid}", "STREAM", "Sự kiện phiên", "XADD MAXLEN ~5000", "7 ngày"),
        ("rankrush:admin:events", "STREAM", "Audit admin", "XADD/XREVRANGE", "Theo chính sách hiện tại"),
        ("rankrush:session:{sid}:phase-lock", "STRING", "Khóa phase phân tán", "SET PX NX + Lua DEL", "5 giây"),
    ]
    add_table(
        doc,
        ["Key pattern", "Type", "Ý nghĩa", "Đọc/Ghi", "Vòng đời"],
        key_rows,
        [2500, 800, 2500, 2100, 1460],
        font_size=7.5,
    )
    add_code_ref(doc, "apps/api/src/store.ts", "L50-L85", "định nghĩa tập trung toàn bộ key")

    add_heading(doc, "4.3 Tư duy thiết kế dữ liệu", 2)
    design_rows = [
        ("Tách entity Hash", "User/quiz/question/session/player là HASH", "Đọc/sửa một field nhanh; tránh serialize toàn document", "Cần giữ index phụ đồng bộ"),
        ("Index thủ công", "email/username/PIN → id bằng STRING", "Lookup O(1)", "Phải cập nhật cùng MULTI"),
        ("Quan hệ bằng SET/LIST", "SET cho membership; LIST cho thứ tự câu hỏi", "Lệnh tự nhiên, đơn giản", "Không có foreign key"),
        ("Leaderboard ZSET", "Member=playerId, score=điểm", "Tăng điểm/rank/top O(log N)", "Tie-break phụ thuộc Redis/member"),
        ("Event Stream", "XADD gần đúng MAXLEN 5000", "Có thứ tự, audit nhanh, giới hạn bộ nhớ", "Không thay thế log bất biến dài hạn"),
        ("Hash tag {sid}", "Các key phiên dùng cùng {sid}", "Chuẩn bị cho Redis Cluster/Lua cùng slot", "Key dài hơn nhưng có chủ đích"),
        ("TTL phiên", "EXPIRE khi tạo/join/submit/update", "Dữ liệu tạm tự dọn", "Báo cáo quá hạn sẽ biến mất"),
        ("Lua atomic", "Join và submit chạy một script", "Không race condition/ghi nửa chừng", "Khó debug hơn code TS"),
    ]
    add_table(doc, ["Quyết định", "Cách làm", "Lợi ích", "Ảnh hưởng/rủi ro"], design_rows,
              [1650, 2350, 2680, 2680], font_size=8.1)

    add_heading(doc, "4.4 Ví dụ ghi và đọc", 2)
    add_para(doc, "Ghi user: kiểm tra index email/username, rồi MULTI ghi Hash + hai index + SET users.")
    add_code_block(
        doc,
        """redis.multi()
  .hset(keys.user(user.id), encode(user))
  .set(keys.userEmail(user.email), user.id)
  .sadd(keys.users, user.id)
  .set(keys.userUsername(username), user.id)
  .exec();""",
    )
    add_code_ref(doc, "apps/api/src/store.ts", "L87-L147")
    add_para(doc, "Đọc user theo email: GET index để lấy id, sau đó HGETALL hash. Đọc nhiều quiz/player dùng pipeline để giảm round-trip.")
    add_code_ref(doc, "apps/api/src/store.ts", "L116-L155, L183-L188, L473-L480")
    add_para(doc, "Ghi điểm: Lua kiểm tra answer đã tồn tại, HSET câu trả lời, SADD index, ZINCRBY điểm, XADD event và EXPIRE trong một thao tác nguyên tử.")
    add_code_ref(doc, "apps/api/src/store.ts", "L541-L599")

    add_heading(doc, "4.5 Redis có gì đặc biệt trong dự án này?", 2)
    for text in [
        "ZADD tạo thành viên leaderboard ban đầu; ZINCRBY tăng điểm không cần đọc-sửa-ghi; ZREVRANGE lấy top; ZREVRANK lấy hạng từ cao xuống; ZRANK cung cấp hạng tăng dần.",
        "SET EX NX vừa tạo TTL vừa khóa/đặt chỗ duy nhất: PIN, cooldown email, phase lock.",
        "XADD MAXLEN ~5000 giữ nhật ký có thứ tự nhưng kiểm soát bộ nhớ.",
        "MULTI gom nhiều lệnh; pipeline tối ưu round-trip; Lua bảo đảm tính nguyên tử xuyên nhiều cấu trúc.",
        "TTL là một phần của mô hình dữ liệu, không cần cron dọn phiên.",
    ]:
        add_bullet(doc, text)
    add_code_ref(doc, "apps/api/src/store.ts", "L369-L430", "Lua join: ZADD, XADD, EXPIRE")
    add_code_ref(doc, "apps/api/src/store.ts", "L483-L538", "ZREVRANGE, ZSCORE, ZREVRANK, ZRANK")
    add_code_ref(doc, "apps/api/src/store.ts", "L541-L599", "ZINCRBY atomic")

    # 5
    doc.add_page_break()
    add_heading(doc, "5. TTL — thời gian sống và vòng đời phiên", 1)
    add_table(
        doc,
        ["Đối tượng", "Biến/cố định", "Mặc định", "Có điều chỉnh?", "Vị trí"],
        [
            ["Session + player + answer + leaderboard + event", "SESSION_TTL_SECONDS", "604800 giây = 7 ngày", "Có, sửa .env rồi restart API", "config.ts L24; store.ts L306-L307, L345, L381, L552"],
            ["PIN tham gia", "JOIN_CODE_TTL_SECONDS", "86400 giây = 1 ngày", "Có", "config.ts L25; store.ts L269-L279, L348"],
            ["Mã đăng ký", "EMAIL_CODE_TTL_SECONDS", "600 giây = 10 phút", "Có", "config.ts L28; index.ts L663-L671"],
            ["Mã đặt lại mật khẩu", "RESET_CODE_TTL_SECONDS", "900 giây = 15 phút", "Có", "config.ts L27; index.ts L663-L671"],
            ["Cooldown gửi OTP", "Cố định", "60 giây", "Chưa có env; cần sửa code", "index.ts L651-L661"],
            ["Login sai", "Cố định", "10 lần / 15 phút", "Chưa có env; cần sửa code", "index.ts L821-L836"],
            ["Phase lock", "Cố định", "5000 ms", "Chưa có env", "index.ts L429-L450"],
            ["Phòng rỗng", "Cố định", "5000 ms grace", "Chưa có env", "index.ts L371-L426"],
            ["JWT host/admin", "Cố định", "7 ngày", "Chưa có env", "auth.ts L21-L22"],
            ["JWT player", "Cố định", "24 giờ", "Chưa có env", "auth.ts L23-L26"],
            ["Countdown/preview/result", "Cố định", "4s / 5s / 5s", "Chưa có env", "index.ts L371-L375"],
        ],
        [1950, 1800, 1800, 1950, 1860],
        font_size=7.8,
    )
    add_callout(
        doc,
        "Trả lời câu hỏi “từ đầu đã có giới hạn không?”",
        "Có. Phiên không sống vĩnh viễn: ngay lúc createSession, HASH session và Stream event đã được EXPIRE 7 ngày. "
        "Mỗi update/join/submit làm mới TTL cho các key đang tồn tại. PIN có TTL độc lập 1 ngày và bị xóa ngay khi ENDED/CANCELLED. "
        "Các user/quiz/question hiện không đặt TTL vì được coi là dữ liệu dài hạn.",
        fill="FFF7E6",
        accent=GOLD,
    )
    add_heading(doc, "5.1 Cách điều chỉnh an toàn", 2)
    add_number(doc, "Sao lưu Redis trước khi đổi vòng đời.")
    add_number(doc, "Đặt SESSION_TTL_SECONDS, JOIN_CODE_TTL_SECONDS, RESET_CODE_TTL_SECONDS, EMAIL_CODE_TTL_SECONDS trong .env.")
    add_number(doc, "Restart API; giá trị mới áp dụng khi key được tạo hoặc được refresh TTL.")
    add_number(doc, "Kiểm tra bằng RedisInsight hoặc redis-cli TTL <key>.")
    add_number(doc, "Nếu cần lưu báo cáo lâu hơn session, tách REPORT_TTL_SECONDS hoặc kho lưu bền vững; không chỉ tăng TTL vô hạn.")
    add_code_ref(doc, ".env.example", "L7-L17")

    # 6
    doc.add_page_break()
    add_heading(doc, "6. Bảo mật, xác thực và cách biết đúng người dùng", 1)
    add_heading(doc, "6.1 Luồng đăng ký chi tiết", 2)
    register_steps = [
        ("1. Form", "React thu displayName, username, email, password, confirmPassword, verificationCode.", "App.tsx L1345-L1410, L1517-L1674"),
        ("2. Gửi OTP", "POST /auth/email-verification/request; Zod kiểm tra email/username.", "App.tsx L1370-L1409; index.ts L715-L754"),
        ("3. Kiểm tra trùng", "GET index email/username; từ chối 409 nếu đã có.", "index.ts L728-L745; store.ts L120-L147"),
        ("4. Rate limit", "SET cooldown EX 60 NX; không cho gửi lại sớm.", "index.ts L651-L661"),
        ("5. Sinh & lưu mã", "crypto.randomInt tạo 6 số; lưu SHA-256 digest + TTL, không lưu OTP rõ.", "index.ts L646-L672"),
        ("6. Gửi mail", "Nodemailer dùng SMTP; nếu không gửi và không cho dev code thì trả lỗi 503.", "mailer.ts L4-L58; index.ts L673-L695"),
        ("7. Submit đăng ký", "POST /auth/register; Zod kiểm tra độ dài/regex/mật khẩu khớp.", "App.tsx L1411-L1430; index.ts L628-L644, L756-L779"),
        ("8. Xác minh OTP", "Băm lại và so digest; tối đa 5 lần sai rồi xóa mã.", "index.ts L761-L778"),
        ("9. Tạo user", "nanoid id; bcrypt cost 12; role HOST/status ACTIVE.", "index.ts L780-L790"),
        ("10. Persist", "MULTI ghi hash user, index email/username, SET users; xóa OTP/attempt.", "store.ts L87-L114; index.ts L791-L794"),
        ("11. Cấp danh tính", "JWT host 7 ngày trả về; client lưu rr_host_token và rr_user.", "index.ts L795-L805; auth.ts L21-L22; App.tsx L1431-L1444"),
    ]
    add_table(doc, ["Bước", "Xử lý", "Truy vết code"], register_steps, [1250, 5100, 3010], font_size=8.2)

    add_heading(doc, "6.2 Luồng đăng nhập sau khi bấm nút", 2)
    login_steps = [
        ("1", "submit() chặn reload, bật busy, POST identifier + password.", "App.tsx L1411-L1430"),
        ("2", "request() thêm Content-Type; nếu có token thì Bearer; parse ApiError.", "api.ts L135-L165"),
        ("3", "API Zod yêu cầu identifier/email và password.", "index.ts L808-L820"),
        ("4", "Đọc rankrush:login-attempts:<identifier>; chặn khi >=10.", "index.ts L821-L830"),
        ("5", "Nếu identifier có @: GET email-index; nếu không: username-index; rồi HGETALL user.", "index.ts L831-L833; store.ts L116-L147"),
        ("6", "bcrypt.compare mật khẩu nhập với passwordHash. Lỗi thì INCR, lần đầu EXPIRE 15 phút.", "index.ts L834-L845"),
        ("7", "Kiểm tra status SUSPENDED; nếu hợp lệ DEL attempts.", "index.ts L847-L854"),
        ("8", "signHost ký JWT với sub/email/role, expiresIn 7d.", "index.ts L855-L865; auth.ts L21-L28"),
        ("9", "Client lưu localStorage rr_host_token + rr_user.", "App.tsx L1431-L1435"),
        ("10", "Nếu có pending quiz thì tạo phòng; ngược lại ADMIN → /admin, HOST → /.", "App.tsx L1436-L1444"),
        ("11", "Các request sau gửi Bearer; requireHost/requireAdmin verify chữ ký, expiry, kind và role.", "api.ts L141-L146; auth.ts L27-L78"),
        ("12", "Socket.IO cũng verify chính JWT ở handshake trước khi join room.", "index.ts L2864-L2907"),
    ]
    add_table(doc, ["#", "Luồng xử lý", "Truy vết code"], login_steps, [600, 5670, 3090], font_size=8.2)
    add_callout(
        doc,
        "Làm sao biết user đó đúng?",
        "Hệ thống không chỉ tin localStorage. Server tra user theo index duy nhất, so mật khẩu bằng bcrypt, kiểm tra trạng thái, "
        "sau đó ký JWT bằng JWT_SECRET. Mỗi endpoint bảo vệ và Socket handshake đều verify chữ ký + thời hạn + loại claims. "
        "Quyền sở hữu quiz/session còn so ownerId/hostId với JWT sub; admin được kiểm tra role.",
        fill="EAF6F0",
        accent=GREEN,
    )
    add_code_ref(doc, "apps/api/src/auth.ts", "L5-L93")
    add_code_ref(doc, "apps/api/src/index.ts", "L176-L230", "kiểm tra owner/host/admin")

    add_heading(doc, "6.3 Quên mật khẩu", 2)
    for text in [
        "Client bước request gửi email; API luôn trả thông điệp trung tính để hạn chế dò email tồn tại.",
        "Nếu user có thật, dùng chung issueEmailCode với purpose=reset; digest TTL mặc định 15 phút.",
        "Submit code + mật khẩu mới; tối đa 5 lần sai; bcrypt cost 12; HDEL rawPassword và DEL mã.",
    ]:
        add_bullet(doc, text)
    add_code_ref(doc, "apps/web/src/App.tsx", "L1713-L1938")
    add_code_ref(doc, "apps/api/src/index.ts", "L868-L933")

    add_heading(doc, "6.4 Điểm mạnh và điểm cần cải thiện bảo mật", 2)
    add_table(
        doc,
        ["Đang có", "Ý nghĩa", "Cần nâng cấp"],
        [
            ["bcrypt cost 12; OTP digest; TTL/rate limit", "Giảm lộ mật khẩu/OTP và brute-force", "Thêm policy mật khẩu, breach check, MFA tùy nhu cầu"],
            ["JWT có expiry và claims", "API/socket stateless", "Đưa thời hạn vào env; refresh token/deny-list khi cần revoke"],
            ["Helmet, CORS private network, Zod", "Giảm lỗi header/origin/input", "Production chỉ allow domain cụ thể; CSP và reverse proxy TLS"],
            ["Redis không có password/TLS trong compose dev", "Dễ chạy local", "Production bật ACL, TLS, bind/private network; không public 6379"],
            ["rawPassword là field legacy tùy chọn", "Admin UI cũ có thể từng cần xem", "Không bao giờ lưu mật khẩu rõ; migration xóa toàn bộ rawPassword"],
            ["SMTP/Gemini secret trong .env", "Tách khỏi code", "Không commit .env; rotate key; secret manager khi deploy"],
        ],
        [2700, 2800, 3860],
        font_size=8.1,
    )
    add_code_ref(doc, "apps/api/src/index.ts", "L99-L169", "CORS, Helmet, giới hạn JSON/upload")
    add_code_ref(doc, "apps/api/src/types.ts", "L21-L32", "User có rawPassword legacy — rủi ro cần loại bỏ")

    # 7
    add_heading(doc, "7. Luồng chi tiết từng nhóm chức năng", 1)
    feature_flows = [
        ("Hồ sơ/avatar", "GET /auth/me → verify host → HGETALL user. PUT /auth/me kiểm tra current password khi đổi password, validate avatar, cập nhật hash/index.", "App.tsx SettingsPage L2644-L2986; index.ts L935-L1139; avatar.ts L1-L58"),
        ("Thư viện/tìm quiz", "GET /quizzes q/category/offset/limit → public ZSET hoặc SET owner → pipeline HGETALL → filter/slice → đếm câu hỏi.", "App.tsx Home/Dashboard L892-L1238, L2051-L2326; index.ts L1705-L1728; store.ts L174-L222"),
        ("Tạo quiz thủ công", "/editor/new chọn type → POST quiz DRAFT → POST blank question → điều hướng editor.", "QuestionTypePicker.tsx L44-L116; index.ts L1819-L1841, L2183-L2196"),
        ("Editor", "GET quiz?editor=1 → owner guard → sửa metadata/câu hỏi → PUT; publish chỉ khi validate được.", "App.tsx L3488-L4260; index.ts L1729-L1763, L1843-L1854, L2183-L2239"),
        ("Clone/xóa", "Clone sao metadata và toàn bộ câu hỏi sang owner; delete xóa question hashes/list/index/quiz hash.", "index.ts L1855-L1890; store.ts L224-L244"),
        ("Public/practice", "Người xem chỉ nhận câu hỏi đã che đáp án; practice gửi answer để server kiểm tra và trả giải thích.", "App.tsx L6482-L6902; index.ts L1729-L1818"),
        ("Tạo phòng", "Host POST /sessions với quizId/settings → validate publishable → tạo questionOrder → reserve PIN SET EX NX → HSET session + Stream + TTL.", "App.tsx createHostRoom L343-L348; index.ts L2243-L2291; store.ts L269-L314"),
        ("Tra PIN/QR", "GET /sessions/pin/:pin lấy session; QR/link dùng origin mạng. PIN hết TTL hoặc đã end trả 404.", "index.ts L596-L625, L2336-L2365; store.ts L319-L322"),
        ("Join player", "Rate-limit IP → tra PIN → ngăn host join phòng mình → lọc nickname → Lua kiểm tra LOBBY/full/trùng, tạo player, ZADD 0, XADD, TTL → JWT player.", "App.tsx L4262-L4448; index.ts L2416-L2463; store.ts L369-L452"),
        ("Lobby", "Socket JWT join room; player online=true; host/player nhận snapshot riêng; host có thể đổi settings hoặc kick khi LOBBY.", "index.ts L2318-L2334, L2700-L2716, L2864-L2908"),
        ("Start game", "Host ownership → state phải LOBBY → phải có socket player online → state GAME_COUNTDOWN → emit + schedule.", "index.ts L2465-L2496"),
        ("Máy trạng thái", "4s countdown → 5s preview → RUNNING theo timeLimit → 5s result → câu tiếp theo/ENDED. Redis phase lock chống hai tiến trình chuyển pha.", "index.ts L371-L585"),
        ("Pause/resume/skip", "Tính elapsed/remaining; lưu pausedState + pausedRemainingMs; resume tái dựng questionStartedAt và schedule; skip sang câu tiếp.", "index.ts L2516-L2611"),
        ("Trả lời", "Player JWT phải thuộc session; validate q hiện tại và deadline; chấm đúng; server tự tính responseMs/score; Lua idempotent ghi answer và tăng ZSET.", "index.ts L2719-L2797; question-answer.ts L1-L57; score.ts L1-L18; store.ts L541-L599"),
        ("Leaderboard", "ZREVRANGE WITHSCORES + pipeline player metadata; self rank bằng ZSCORE/ZREVRANK; team tương tự.", "store.ts L483-L538; index.ts L2822-L2853"),
        ("Kết thúc/báo cáo", "End thủ công, tự hết câu hoặc phòng rỗng; report tổng hợp answers, accuracy, avgResponseMs; player result chỉ sau ENDED.", "index.ts L396-L426, L545-L564, L2624-L2641, L2799-L2861; store.ts L611-L699"),
        ("Replay/cancel", "Cancel chỉ LOBBY; replay chỉ ENDED và tạo session/PIN mới với settings cũ.", "index.ts L2643-L2698"),
        ("Mất kết nối", "Socket disconnect chỉ set offline khi không còn socket cùng player; sau grace 5s, nếu active game không còn ai thì ENDED và báo host.", "index.ts L2909-L2924; session-presence.ts L3-L20"),
        ("Ngôn ngữ/theme", "Preferences lưu locale/theme localStorage, resolve system theme; header/admin dùng chung context.", "preferences.tsx L208-L296; App.tsx PreferenceControls L570-L595"),
        ("Voice câu hỏi", "Web Speech API ghép prompt + options, chọn vi-VN/en-US, lưu bật/tắt localStorage và tự đọc khi câu đổi.", "QuestionVoice.tsx L6-L78"),
        ("Nhạc/âm thanh", "PlayerMusic quản lý nguồn local/YouTube và lựa chọn người dùng.", "PlayerMusic.tsx L1-L302; background-music.ts L1-L77"),
        ("Copy link/PIN", "copyText ưu tiên Clipboard API, fallback textarea + execCommand.", "App.tsx L316-L334"),
    ]
    add_table(doc, ["Chức năng", "Quy trình thực hiện", "Code"], feature_flows,
              [1500, 5150, 2710], font_size=7.55)

    add_heading(doc, "7.1 Máy trạng thái phiên game", 2)
    add_code_block(
        doc,
        """LOBBY
  └─ start (có player online) → GAME_COUNTDOWN [4s]
       └─ tự động → QUESTION_PREVIEW [5s]
            └─ tự động/open → RUNNING [timeLimitSec]
                 ├─ tất cả trả lời hoặc hết giờ → QUESTION_RESULT [5s]
                 │    ├─ còn câu → QUESTION_PREVIEW
                 │    └─ hết câu → ENDED
                 ├─ pause → PAUSED → resume về phase trước
                 └─ không còn player online sau 5s → ENDED
LOBBY ─ cancel → CANCELLED
ENDED ─ replay → session mới ở LOBBY""",
    )

    add_heading(doc, "7.2 Chấm điểm và chống gửi lặp", 2)
    add_para(
        doc,
        "Server không tin responseMs của client; dùng Date.now() - questionStartedAt. Hàm calculateScore trả 0 nếu sai, "
        "nếu đúng thì base 500 + tối đa 500 theo tốc độ, cộng 200 cho câu cuối. Answer key chứa session/player/question; "
        "Lua EXISTS bảo đảm lần gửi sau không cộng điểm lại.",
    )
    add_code_ref(doc, "apps/api/src/index.ts", "L2756-L2796")
    add_code_ref(doc, "apps/api/src/score.ts", "L1-L18")
    add_code_ref(doc, "apps/api/src/store.ts", "L541-L599")

    # 8
    doc.add_page_break()
    add_heading(doc, "8. Trí tuệ nhân tạo, PDF/CSV và kiểm định câu hỏi", 1)
    add_heading(doc, "8.1 Luồng sinh quiz", 2)
    ai_steps = [
        ("1", "Host chọn SUBJECT/PDF/CSV, nhập title/category/count/language/difficulty/context.", "App.tsx L2988-L3051"),
        ("2", "Frontend FormData POST /ai/generate-quiz với Bearer.", "App.tsx L3062-L3105"),
        ("3", "Multer memoryStorage giới hạn 10MB, chỉ PDF/CSV.", "index.ts L148-L169"),
        ("4", "CSV: parser kiểm tra template và map trực tiếp sang Question.", "index.ts L1907-L2181; csv-quiz-import.ts L1-L220"),
        ("5", "PDF: pdf-parse trích text; PDF scan ảnh không có OCR sẽ có thể rỗng.", "index.ts L1907-L2181"),
        ("6", "Router ưu tiên AI_PROVIDER; hiện cấu hình Gemini, fallback Ollama, sau cùng local generator.", "config.ts L42-L59; ollama-quiz-generator.ts L214-L283"),
        ("7", "Gemini yêu cầu structured JSON schema, blueprint đa dạng type và ngôn ngữ.", "gemini-quiz-generator.ts L14-L203, L234-L353"),
        ("8", "Nếu review bật: chấm quality, phát hiện câu yếu, tạo lại phần không đạt ngưỡng.", "gemini-quiz-generator.ts L384-L526"),
        ("9", "Question diversity loại trùng/lặp gần.", "question-diversity.ts L1-L51"),
        ("10", "API tạo quiz DRAFT + save từng question; client chuyển sang editor để người dùng kiểm tra.", "index.ts L1907-L2181; App.tsx L3106-L3127"),
    ]
    add_table(doc, ["#", "Luồng", "Code"], ai_steps, [600, 5680, 3080], font_size=8.2)
    add_heading(doc, "8.2 Mô hình và trade-off", 2)
    add_table(
        doc,
        ["Provider", "Vai trò", "Ưu điểm", "Nhược điểm/rủi ro"],
        [
            ["Gemini (gemini-3.6-flash theo .env.example)", "Primary + reviewer", "Structured output, ngữ cảnh tốt, chất lượng cao hơn local nhỏ", "API key, quota/chi phí, mạng, model name có thể thay đổi"],
            ["Ollama qwen2.5:3b", "Fallback local", "Riêng tư, offline, chủ động", "3B dễ rập khuôn/hallucination; máy yếu chậm"],
            ["Local generator", "Fallback cuối", "Luôn có kết quả, không phụ thuộc dịch vụ", "Template-based, ít sáng tạo, chỉ phù hợp dự phòng"],
            ["CSV", "Đường import deterministic", "Kiểm soát nội dung tuyệt đối, dễ kiểm thử", "Người dùng phải chuẩn bị dữ liệu đúng mẫu"],
        ],
        [1800, 1900, 2700, 2960],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Nguyên tắc chất lượng",
        "AI chỉ tạo bản nháp. Câu hỏi và đáp án đúng phải được host kiểm tra trước khi publish. "
        "Đối với PDF học thuật, nên bổ sung trích dẫn đoạn nguồn/page, OCR cho PDF scan và bộ đánh giá fact-grounding.",
        fill="FFF7E6",
        accent=GOLD,
    )
    add_code_ref(doc, "apps/api/src/gemini-quiz-generator.ts", "L466-L562")
    add_code_ref(doc, "apps/api/src/ollama-quiz-generator.ts", "L99-L283")
    add_code_ref(doc, "apps/api/src/quiz-generator.ts", "L387-L430")

    # 9
    add_heading(doc, "9. Quản trị, sao lưu, giám sát và vận hành", 1)
    admin_rows = [
        ("Tổng quan", "Đếm users/sessions/quizzes, trạng thái và hoạt động gần đây.", "AdminPage.tsx L757-L880; index.ts L1142-L1195"),
        ("Người dùng", "Tìm/filter/phân trang; tạo; role/status; reset mật khẩu; bảo vệ admin cuối và self.", "AdminPage.tsx L882-L1278; index.ts L1197-L1449"),
        ("Phòng chơi", "Filter trạng thái, tìm, phân trang, ẩn/hiện PIN, mở/kết thúc phòng.", "AdminPage.tsx L1280-L1508; index.ts L1451-L1463"),
        ("Nhật ký", "Đọc Stream admin + events; filter/tìm/phân trang phía UI/API.", "AdminPage.tsx L1510-L1651; index.ts L1465-L1527"),
        ("Hệ thống Redis", "INFO server/memory/clients/stats, keyspace/dbsize.", "AdminPage.tsx L1662-L1997; index.ts L1529-L1566"),
        ("Backup", "SCAN namespace, DUMP/PTTL, checksum; file JSON; restore có safety backup và rollback.", "backup-service.ts L73-L347; index.ts L1568-L1680"),
    ]
    add_table(doc, ["Trang", "Xử lý", "Code"], admin_rows, [1500, 5050, 2810], font_size=8.25)
    add_heading(doc, "9.1 Redis persistence và công cụ vận hành", 2)
    for text in [
        "Docker chạy redis:7.4-alpine với appendonly yes và appendfsync everysec; dữ liệu gắn named volume rankrush_redis_data.",
        "AOF everysec cân bằng bền vững/hiệu năng nhưng có thể mất khoảng một giây dữ liệu khi sự cố đột ngột.",
        "Healthcheck PING bảo đảm Redis khỏe trước RedisInsight.",
        "RedisInsight 2.70.1 publish 5540:5540 để xem key, type, TTL và query trực quan.",
        "Backup ứng dụng có checksum, namespace validation, khóa thao tác, safety backup và rollback; không thay cho snapshot/offsite/DR production.",
    ]:
        add_bullet(doc, text)
    add_code_ref(doc, "docker-compose.yml", "L1-L28")
    add_code_ref(doc, "apps/api/src/backup-service.ts", "L17-L347")

    add_heading(doc, "9.2 Cách chạy và kiểm tra", 2)
    add_code_block(
        doc,
        """docker compose up -d
npm install
npm run dev

# Web:          http://localhost:5173
# API health:   http://localhost:4000/api/health
# RedisInsight: http://localhost:5540
# Redis URL:    redis://localhost:6379""",
    )
    add_para(
        doc,
        "Nếu API báo EADDRINUSE :4000, một tiến trình API khác đang chạy. Xác định PID bằng "
        "Get-NetTCPConnection -LocalPort 4000, dừng đúng tiến trình hoặc đổi PORT và cấu hình proxy tương ứng. "
        "Không khởi động hai API trên cùng cổng.",
    )
    add_code_ref(doc, "package.json", "L8-L16")
    add_code_ref(doc, "apps/api/src/index.ts", "L2977-L2987")

    # 10
    add_heading(doc, "10. Tối ưu, giới hạn và lộ trình nâng cấp", 1)
    add_heading(doc, "10.1 Những tối ưu đã có", 2)
    for text in [
        "Pipeline khi đọc nhiều hash; MULTI khi cập nhật index liên quan.",
        "ZSET cho top/rank thay vì sort toàn bộ trong Node.",
        "Lua cho join/answer idempotent và nguyên tử.",
        "Hash tag {sessionId} giúp các key cùng phiên ở cùng slot khi chuyển Redis Cluster.",
        "TTL và MAXLEN kiểm soát dữ liệu tạm/Stream.",
        "Snapshot theo viewer che đáp án/rank và giảm rò rỉ dữ liệu.",
        "Zod, upload limit, rate limit, nickname filter và owner guards tại biên.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "10.2 Điểm nghẽn/rủi ro", 2)
    add_table(
        doc,
        ["Vấn đề", "Ảnh hưởng", "Khuyến nghị"],
        [
            ["Redis là kho chính in-memory", "Chi phí RAM; dữ liệu lịch sử phụ thuộc TTL/AOF", "Tách dữ liệu nóng/lạnh; archive report sang PostgreSQL/object storage"],
            ["listUsers/listSessions đọc toàn SET rồi pipeline", "Chậm khi hàng trăm nghìn bản ghi", "Index phân trang bằng ZSET theo createdAt; cursor pagination"],
            ["Filter quiz ở Node", "Tốn CPU/RAM khi thư viện lớn", "RediSearch hoặc PostgreSQL full-text/search service"],
            ["countAnswersForQuestion quét mọi answer key", "O(số câu trả lời) mỗi lần submit", "Dùng counter/hash per question hoặc bitmap/set riêng"],
            ["Socket.IO chỉ một instance", "Scale ngang mất room/presence chung", "Socket.IO Redis adapter, sticky session, distributed timers/worker"],
            ["Timer nằm trong RAM process", "Restart API làm mất timer đang chạy", "Lưu deadline trong Redis + scheduler/worker phục hồi"],
            ["JWT không revoke", "Token bị lộ còn hiệu lực", "Refresh token rotation, tokenVersion hoặc deny-list TTL"],
            ["Admin raw password concept", "Rủi ro nghiêm trọng nếu giữ mật khẩu rõ", "Xóa rawPassword hoàn toàn; chỉ reset, không hiển thị mật khẩu"],
            ["Redis port 6379 publish local", "Nguy hiểm nếu bind ra mạng công cộng", "Firewall/private network/ACL/TLS; không public production"],
            ["AOF everysec đơn nút", "Không HA; có RPO ~1s", "Replica + Sentinel/managed Redis; backup offsite và diễn tập restore"],
        ],
        [2500, 2850, 4010],
        font_size=8.0,
    )

    add_heading(doc, "10.3 Lộ trình nâng cấp đề xuất", 2)
    roadmap = [
        ("P0 — Bảo mật", "Xóa rawPassword; rotate secrets; Redis ACL/TLS/private; production CORS; backup offsite.", "Giảm rủi ro mất dữ liệu/tài khoản"),
        ("P1 — Độ tin cậy", "Timer phục hồi từ deadline Redis; Redis replica/Sentinel; metrics/log tập trung.", "Game tiếp tục đúng khi restart/lỗi node"),
        ("P1 — Hiệu năng", "Counter answer theo question; ZSET pagination users/sessions; cache snapshot ngắn.", "Giảm scan và round-trip"),
        ("P2 — Scale realtime", "Socket.IO Redis adapter + worker phase + idempotency keys.", "Nhiều API instance"),
        ("P2 — Dữ liệu dài hạn", "Outbox/Stream consumer ghi report/audit sang PostgreSQL hoặc warehouse.", "Phân tích, truy vấn và lưu lâu"),
        ("P3 — AI", "RAG theo trang PDF, OCR, citations, evaluation dataset, A/B model và cost telemetry.", "Câu hỏi sát nguồn và ít rập khuôn"),
    ]
    add_table(doc, ["Ưu tiên", "Công việc", "Kết quả"], roadmap, [1800, 5210, 2350], font_size=8.3)

    add_heading(doc, "10.4 Có thể nâng cấp CSDL không?", 2)
    add_para(
        doc,
        "Có. Thiết kế key dùng prefix tập trung và hash tag session đã tạo nền tảng tốt. Nâng cấp ít rủi ro nhất là: "
        "(1) thêm index mới song song, backfill bằng pipeline; (2) dual-write/Stream consumer cho kho bền vững; "
        "(3) chuyển Redis standalone sang managed Redis/replica/Sentinel rồi Cluster khi cần; "
        "(4) version hóa key/payload; (5) giữ API contract ổn định. Không nên đổi toàn bộ sang MongoDB chỉ vì dữ liệu là JSON; "
        "cần xuất phát từ truy vấn, độ bền, quy mô và SLA thực tế.",
    )

    # 11
    doc.add_page_break()
    add_heading(doc, "11. Danh mục API và chỉ mục mã nguồn", 1)
    api_rows = [
        ("GET", "/api/health", "Ping Redis", "index.ts L589-L595"),
        ("GET", "/api/meta/network", "Origin LAN/QR", "index.ts L596-L625"),
        ("GET/POST", "/api/auth/email-status; /email-verification/request", "SMTP/OTP đăng ký", "index.ts L699-L754"),
        ("POST", "/api/auth/register", "Đăng ký", "index.ts L756-L807"),
        ("POST", "/api/auth/login", "Đăng nhập", "index.ts L808-L867"),
        ("POST", "/api/auth/forgot-password; /reset-password", "Khôi phục mật khẩu", "index.ts L868-L934"),
        ("GET/PUT", "/api/auth/me", "Hồ sơ", "index.ts L935-L1139"),
        ("GET", "/api/admin/overview", "Tổng quan admin", "index.ts L1142-L1195"),
        ("GET/POST/PATCH/PUT", "/api/admin/users...", "Quản trị user", "index.ts L1197-L1449"),
        ("GET", "/api/admin/sessions; /activity; /system", "Phòng/log/Redis", "index.ts L1451-L1566"),
        ("GET/POST/DELETE", "/api/admin/backups...", "Backup/restore/download", "index.ts L1568-L1680"),
        ("GET/POST/PUT/DELETE", "/api/quizzes...", "Quiz CRUD/clone/practice", "index.ts L1705-L1890"),
        ("GET/POST", "/api/ai/status; /csv-template; /generate-quiz", "AI/import", "index.ts L1892-L2181"),
        ("POST/PUT/DELETE", "/api/quizzes/:id/questions; /api/questions/:id", "Question CRUD", "index.ts L2183-L2239"),
        ("GET/POST/PATCH", "/api/sessions...", "Tạo/list/snapshot/settings/join", "index.ts L2243-L2463"),
        ("POST", "/api/sessions/:id/start|advance|skip|pause|resume|open-question", "Điều khiển phase", "index.ts L2465-L2623"),
        ("POST", "/api/sessions/:id/end|cancel|replay|kick", "Vòng đời phòng", "index.ts L2624-L2717"),
        ("POST/GET", "/api/sessions/:id/answers|result|leaderboard|report", "Trả lời/kết quả", "index.ts L2719-L2861"),
        ("Socket", "session:join + disconnect", "Realtime/presence", "index.ts L2864-L2927"),
    ]
    add_table(doc, ["Method", "Endpoint", "Mục đích", "Code"], api_rows,
              [900, 3300, 2800, 2360], font_size=7.8)

    add_heading(doc, "11.1 Chỉ mục tệp quan trọng", 2)
    index_rows = [
        ("apps/api/src/index.ts", "HTTP routes, socket, game state machine, auth flow, AI upload"),
        ("apps/api/src/store.ts", "Mô hình key, CRUD Redis, ZSET, Lua join/submit, report"),
        ("apps/api/src/config.ts", "Đọc .env, TTL, Redis, SMTP, AI"),
        ("apps/api/src/redis.ts", "Kết nối/đóng Redis"),
        ("apps/api/src/auth.ts", "JWT và middleware quyền"),
        ("apps/api/src/mailer.ts", "SMTP và email OTP"),
        ("apps/api/src/types.ts", "Mô hình domain"),
        ("apps/api/src/gemini-quiz-generator.ts", "Sinh/review Gemini"),
        ("apps/api/src/ollama-quiz-generator.ts", "Ollama và provider router"),
        ("apps/api/src/quiz-generator.ts", "Local fallback"),
        ("apps/api/src/backup-service.ts", "Backup/restore/checksum/rollback"),
        ("apps/web/src/App.tsx", "Các trang, form, editor, game host/player/report"),
        ("apps/web/src/api.ts", "HTTP client, token, lỗi bản địa hóa"),
        ("apps/web/src/AdminPage.tsx", "Admin console"),
        ("apps/web/src/preferences.tsx", "Theme/locale"),
        ("apps/web/src/QuestionVoice.tsx", "Web Speech API"),
        ("docker-compose.yml", "Redis/AOF/volume/RedisInsight"),
        (".env.example", "Mẫu biến runtime không chứa secret thật"),
    ]
    add_table(doc, ["Tệp", "Vai trò"], index_rows, [3800, 5560], font_size=8.6)

    add_heading(doc, "11.2 Các truy vấn Redis minh họa khi thuyết trình", 2)
    add_code_block(
        doc,
        """# Kiểm tra kết nối
PING

# Xem kiểu và TTL phiên
TYPE rankrush:session:{<sessionId>}
TTL  rankrush:session:{<sessionId>}

# Top 10 và rank
ZREVRANGE rankrush:leaderboard:{<sessionId>} 0 9 WITHSCORES
ZREVRANK  rankrush:leaderboard:{<sessionId>} <playerId>
ZRANK     rankrush:leaderboard:{<sessionId>} <playerId>

# Xem entity/index
HGETALL rankrush:user:<userId>
GET     rankrush:user-email:<email>
LRANGE  rankrush:quiz:<quizId>:questions 0 -1
XRANGE  rankrush:events:{<sessionId>} - + COUNT 20""",
    )
    add_para(
        doc,
        "Khi demo, thay placeholder bằng id thực và không chiếu passwordHash, SMTP_PASS, JWT_SECRET hoặc API key. "
        "RedisInsight nên kết nối host.docker.internal:6379 (nếu RedisInsight trong container và Redis publish trên host) "
        "hoặc hostname dịch vụ redis:6379 khi cùng Compose network.",
        size=9.5,
        color=MUTED,
    )

    add_heading(doc, "12. Kết luận", 1)
    add_para(
        doc,
        "RankRush đã khai thác đúng các thế mạnh khác biệt của Redis: ZSET cho leaderboard, TTL cho dữ liệu tạm, "
        "Lua cho cập nhật nguyên tử, Stream cho sự kiện và pipeline/MULTI để giảm round-trip. Luồng đăng ký/đăng nhập "
        "có kiểm tra runtime, OTP digest, bcrypt, JWT và owner guards nên server có căn cứ xác định người dùng thay vì "
        "tin giao diện. Kiến trúc hiện tại phù hợp bài toán game realtime và trình diễn học thuật; lộ trình production "
        "nên ưu tiên loại bỏ mọi khả năng lưu mật khẩu rõ, tăng độ bền/HA, phục hồi timer và tách dữ liệu lịch sử.",
    )
    add_callout(
        doc,
        "Câu trả lời ngắn gọn cho hội đồng",
        "Dự án chọn Redis vì dữ liệu game là dữ liệu nóng, truy cập theo khóa, cần cập nhật điểm/rank tức thời và tự hết hạn. "
        "Thiết kế tách Hash/Set/List/ZSET/Stream theo đúng hình dạng truy vấn; Lua bảo vệ tính nguyên tử. "
        "MongoDB không được sử dụng trong phiên bản này. Redis rất mạnh cho realtime, nhưng dữ liệu phân tích dài hạn nên "
        "được sao lưu hoặc đồng bộ sang một kho bền vững khi hệ thống mở rộng.",
        fill="EAF6F0",
        accent=GREEN,
    )

    # Core properties and document hygiene.
    doc.core_properties.title = "Báo cáo kỹ thuật và luồng xử lý RankRush"
    doc.core_properties.subject = "Công nghệ, Redis, TTL, xác thực và toàn bộ chức năng"
    doc.core_properties.author = "Codex — phân tích từ mã nguồn RankRush"
    doc.core_properties.keywords = "RankRush, Redis, TypeScript, React, Node.js, Socket.IO, TTL, ZSET"
    doc.core_properties.comments = "Nguồn: nhánh develop, commit 5932792, rà soát 02/08/2026."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
